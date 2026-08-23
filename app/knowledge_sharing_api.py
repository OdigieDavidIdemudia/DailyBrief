from fastapi import APIRouter, HTTPException, Depends
from app.auth import get_current_user, UserSession
from app.schemas_ks import KSGenerateRequest, KSRefineRequest, KSExportRequest
from app.knowledge_sharing_ai import generate_ks_draft, refine_ks_draft
from fastapi.responses import FileResponse
import os, json, uuid

router = APIRouter()

@router.post("/api/ks/generate")
async def api_ks_generate(req: KSGenerateRequest, current_user: UserSession = Depends(get_current_user)):
    try:
        draft = await generate_ks_draft(req)
        if "error" in draft:
            raise HTTPException(status_code=500, detail=draft["error"])
        return {"success": True, "draft": draft}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/api/ks/refine")
async def api_ks_refine(req: KSRefineRequest, current_user: UserSession = Depends(get_current_user)):
    try:
        draft = await refine_ks_draft(req.draft, req.corrections)
        if "error" in draft:
            raise HTTPException(status_code=500, detail=draft["error"])
        return {"success": True, "draft": draft}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/api/ks/export")
async def api_ks_export(req: KSExportRequest, current_user: UserSession = Depends(get_current_user)):
    try:
        static_dir = os.path.join(os.getcwd(), "static", "reports")
        os.makedirs(static_dir, exist_ok=True)
        title_slug = req.draft.get("presentation_title", "KS")[:30].replace(" ", "_").replace(":", "").replace("&", "and")
        uid = uuid.uuid4().hex[:6]

        if req.format == "txt":
            out_filename = f"{title_slug}_{uid}.txt"
            out_path = os.path.join(static_dir, out_filename)
            lines = []
            lines.append("=" * 70)
            lines.append(req.draft.get("presentation_title", "").upper())
            lines.append(f"Audience: {req.draft.get('target_audience', '')}  |  Author: {req.draft.get('author', '')}  |  Date: {req.draft.get('date', '')}")
            lines.append("=" * 70)
            for slide in req.draft.get("slides", []):
                lines.append(f"\n{'─' * 60}")
                lines.append(f"SLIDE {slide['slide_number']}: {slide['title'].upper()}")
                if slide.get("subtitle"):
                    lines.append(f"  \"{slide['subtitle']}\"")
                lines.append("─" * 60)
                if slide.get("key_definition") or slide.get("definition"):
                    lines.append(f"\n  [DEFINITION]\n  {slide.get('key_definition', slide.get('definition'))}\n")
                lines.append("  KEY POINTS:")
                for b in slide.get("items", []):
                    lines.append(f"    • {b.get('label')}: {b.get('content')}")
                if slide.get("flow_steps"):
                    for st in slide.get("flow_steps", []):
                        lines.append(f"    -> Step {st.get('step_number')}: {st.get('label')} - {st.get('detail')}")
                lines.append(f"\n  [PRESENTER NOTES]\n  {slide.get('presenter_script', '')}")
            lines.append("\n" + "=" * 70 + "\n  END OF GUIDE\n" + "=" * 70)
            with open(out_path, "w", encoding="utf-8") as f:
                f.write("\n".join(lines))
            return FileResponse(out_path, filename=out_filename, media_type="text/plain")

        elif req.format == "pptx":
            from pptx import Presentation
            from pptx.util import Inches, Pt, Emu
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN
            
            def hex2rgb(h):
                h = str(h).lstrip('#')
                if len(h) != 6: return RGBColor(143, 168, 255) # default info accent
                return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

            # Brand colours from Spec
            BG_DARK = RGBColor(0x11, 0x14, 0x1C)
            SURF_1  = RGBColor(0x1B, 0x21, 0x30)
            SURF_2  = RGBColor(0x23, 0x2A, 0x38)
            TEXT_P  = RGBColor(0xF5, 0xF6, 0xF8)
            TEXT_S  = RGBColor(0xB4, 0xBB, 0xC8)
            
            accent_hex = req.draft.get("accent_primary", "8FA8FF")
            ACCENT = hex2rgb(accent_hex)

            W, H = Inches(13.33), Inches(7.5)

            prs = Presentation()
            prs.slide_width = W
            prs.slide_height = H
            blank = prs.slide_layouts[6]

            def add_rect(slide, l, t, w, h, fill_rgb, outline=None):
                shape = slide.shapes.add_shape(1, l, t, w, h)
                shape.fill.solid()
                shape.fill.fore_color.rgb = fill_rgb
                if outline:
                    shape.line.color.rgb = outline
                    shape.line.width = Pt(1)
                else:
                    shape.line.fill.background()
                return shape

            def add_text(slide, l, t, w, h, text, size, bold=False, color=TEXT_P, align=PP_ALIGN.LEFT, italic=False):
                tb = slide.shapes.add_textbox(l, t, w, h)
                tf = tb.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = align
                run = p.add_run()
                run.text = text
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.italic = italic
                run.font.color.rgb = color
                run.font.name = "Arial"
                return tb

            slides_data = req.draft.get("slides", [])
            total = len(slides_data)
            pres_title = req.draft.get("presentation_title", "Knowledge Sharing")

            # ── TITLE SLIDE ───────────────────────────────────────────────────
            sl = prs.slides.add_slide(blank)
            add_rect(sl, 0, 0, W, H, BG_DARK)
            add_rect(sl, 0, 0, Inches(0.2), H, ACCENT)
            
            add_text(sl, Inches(1), Inches(2.2), Inches(11), Inches(1.8), pres_title, 36, bold=True, color=TEXT_P)
            add_rect(sl, Inches(1), Inches(4.2), Inches(6), Inches(0.05), ACCENT)
            
            first_slide = slides_data[0] if slides_data else {}
            hook = first_slide.get("subtitle", f"Audience: {req.draft.get('target_audience','')}")
            add_text(sl, Inches(1), Inches(4.4), Inches(10), Inches(0.5), hook, 16, color=TEXT_S)
            
            meta = f"{req.draft.get('author','')}  |  {req.draft.get('target_audience','')}  |  {req.draft.get('date','')}"
            add_text(sl, Inches(1), H - Inches(1), Inches(10), Inches(0.4), meta, 11, color=TEXT_S)

            # ── CONTENT SLIDES ────────────────────────────────────────────────
            for idx, slide_data in enumerate(slides_data):
                sl = prs.slides.add_slide(blank)
                add_rect(sl, 0, 0, W, H, BG_DARK)

                # Header
                eyebrow = slide_data.get("eyebrow", "")
                if eyebrow:
                    add_text(sl, Inches(0.6), Inches(0.4), Inches(10), Inches(0.3), eyebrow.upper(), 9.5, bold=True, color=ACCENT)
                
                add_text(sl, Inches(0.6), Inches(0.7), Inches(11), Inches(0.6), slide_data.get("title", ""), 30, bold=True)
                
                subtitle = slide_data.get("subtitle", "")
                if subtitle:
                    add_text(sl, Inches(0.6), Inches(1.35), Inches(11), Inches(0.4), subtitle, 14, color=TEXT_S)

                # Bottom Footer
                add_rect(sl, 0, H - Inches(0.5), W, Inches(0.02), SURF_2)
                add_text(sl, Inches(0.5), H - Inches(0.45), Inches(8), Inches(0.3), f"GTCO Security | {pres_title}", 9, color=TEXT_S)
                add_text(sl, W - Inches(1.5), H - Inches(0.45), Inches(1), Inches(0.3), f"{idx+1:02d} / {total:02d}", 10, bold=True, color=TEXT_P, align=PP_ALIGN.RIGHT)

                layout_type = slide_data.get("layout_type", "structured_list")
                items = slide_data.get("items", [])
                steps = slide_data.get("flow_steps", [])
                
                content_top = Inches(2.0)
                
                # MAIN MESSAGE STRIP
                msg = slide_data.get("main_message", "")
                if msg:
                    add_rect(sl, Inches(0.6), content_top, W - Inches(1.2), Inches(0.6), SURF_1, outline=SURF_2)
                    add_rect(sl, Inches(0.6), content_top, Inches(0.1), Inches(0.6), ACCENT)
                    add_text(sl, Inches(0.8), content_top + Inches(0.1), W - Inches(1.6), Inches(0.4), msg, 12.5, bold=True, color=TEXT_P)
                    content_top += Inches(0.8)

                if layout_type == "definition_plus_example":
                    # Definition box
                    add_rect(sl, Inches(0.6), content_top, Inches(6), Inches(1.5), SURF_1)
                    add_text(sl, Inches(0.8), content_top + Inches(0.2), Inches(5.6), Inches(0.3), "DEFINITION", 9.5, bold=True, color=ACCENT)
                    add_text(sl, Inches(0.8), content_top + Inches(0.6), Inches(5.6), Inches(1), slide_data.get("definition", ""), 14, color=TEXT_P)
                    
                    # Examples
                    if items:
                        add_text(sl, Inches(7), content_top + Inches(0.2), Inches(5), Inches(0.3), "KEY ASPECTS", 9.5, bold=True, color=ACCENT)
                        for i, item in enumerate(items[:3]):
                            y = content_top + Inches(0.6) + (i * Inches(0.7))
                            add_text(sl, Inches(7), y, Inches(5), Inches(0.25), item.get("label", ""), 12.5, bold=True, color=TEXT_P)
                            add_text(sl, Inches(7), y + Inches(0.25), Inches(5), Inches(0.4), item.get("content", ""), 11, color=TEXT_S)

                elif layout_type in ["horizontal_process", "attack_chain", "timeline"]:
                    # Horizontal Flow
                    num_steps = len(steps)
                    if num_steps > 0:
                        box_w = min(Inches(11.5) / num_steps, Inches(3))
                        gap = Inches(0.2)
                        start_x = Inches(0.6)
                        for i, st in enumerate(steps):
                            x = start_x + (i * (box_w + gap))
                            add_rect(sl, x, content_top, box_w, Inches(2), SURF_1)
                            add_rect(sl, x, content_top, box_w, Inches(0.05), ACCENT)
                            # Number
                            add_text(sl, x + Inches(0.2), content_top + Inches(0.2), Inches(0.5), Inches(0.3), str(st.get("step_number", i+1)), 18, bold=True, color=ACCENT)
                            add_text(sl, x + Inches(0.2), content_top + Inches(0.6), box_w - Inches(0.4), Inches(0.3), st.get("label", ""), 12.5, bold=True)
                            add_text(sl, x + Inches(0.2), content_top + Inches(0.9), box_w - Inches(0.4), Inches(1), st.get("detail", ""), 11, color=TEXT_S)
                            # Arrow
                            if i < num_steps - 1:
                                add_text(sl, x + box_w, content_top + Inches(0.8), gap, Inches(0.5), "→", 18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
                
                elif layout_type in ["two_column_comparison", "before_after"]:
                    col_w = Inches(5.8)
                    for i, item in enumerate(items[:2]):
                        x = Inches(0.6) if i == 0 else Inches(6.8)
                        add_rect(sl, x, content_top, col_w, Inches(3), SURF_1)
                        add_text(sl, x + Inches(0.4), content_top + Inches(0.4), col_w - Inches(0.8), Inches(0.3), item.get("label", "").upper(), 14, bold=True, color=ACCENT)
                        add_text(sl, x + Inches(0.4), content_top + Inches(0.8), col_w - Inches(0.8), Inches(2), item.get("content", ""), 12.5, color=TEXT_P)
                        if i == 0:
                            add_text(sl, Inches(6.4), content_top + Inches(1.3), Inches(0.4), Inches(0.5), "VS", 14, bold=True, color=TEXT_S, align=PP_ALIGN.CENTER)

                elif layout_type == "risk_control_mapping":
                    risks = [it for it in items if it.get("role") == "risk"]
                    controls = [it for it in items if it.get("role") == "control"]
                    r_w = Inches(5)
                    c_w = Inches(6)
                    add_text(sl, Inches(0.6), content_top, r_w, Inches(0.3), "RISK FACTOR", 9.5, bold=True, color=RGBColor(0xE8, 0x99, 0x8C))
                    add_text(sl, Inches(6.0), content_top, c_w, Inches(0.3), "MITIGATING CONTROL", 9.5, bold=True, color=RGBColor(0x7F, 0xD8, 0xCC))
                    
                    y = content_top + Inches(0.4)
                    for i in range(max(len(risks), len(controls))):
                        if i < len(risks):
                            add_rect(sl, Inches(0.6), y, r_w, Inches(0.8), SURF_1, outline=RGBColor(0xE8, 0x99, 0x8C))
                            add_text(sl, Inches(0.8), y + Inches(0.1), r_w - Inches(0.4), Inches(0.25), risks[i].get("label",""), 11, bold=True)
                            add_text(sl, Inches(0.8), y + Inches(0.35), r_w - Inches(0.4), Inches(0.4), risks[i].get("content",""), 10, color=TEXT_S)
                        if i < len(controls):
                            add_rect(sl, Inches(6.0), y, c_w, Inches(0.8), SURF_1, outline=RGBColor(0x7F, 0xD8, 0xCC))
                            add_text(sl, Inches(6.2), y + Inches(0.1), c_w - Inches(0.4), Inches(0.25), controls[i].get("label",""), 11, bold=True)
                            add_text(sl, Inches(6.2), y + Inches(0.35), c_w - Inches(0.4), Inches(0.4), controls[i].get("content",""), 10, color=TEXT_S)
                        # Connector
                        if i < len(risks) and i < len(controls):
                            add_rect(sl, Inches(5.6), y + Inches(0.4), Inches(0.4), Inches(0.02), TEXT_S)
                        y += Inches(0.9)

                else:
                    # Default: structured_list or similar
                    num_items = len(items)
                    if num_items > 0:
                        per_item = min(Inches(0.8), Inches(3.5) / num_items)
                        for i, item in enumerate(items[:5]):
                            y = content_top + (i * per_item)
                            # Accent dot
                            add_rect(sl, Inches(0.6), y + Inches(0.08), Inches(0.1), Inches(0.1), ACCENT)
                            add_text(sl, Inches(0.8), y, Inches(3.5), per_item, item.get("label", ""), 12.5, bold=True, color=TEXT_P)
                            add_text(sl, Inches(4.5), y, Inches(8.0), per_item, item.get("content", ""), 12.5, color=TEXT_S)

                # Presenter notes
                notes = sl.notes_slide
                notes.notes_text_frame.text = slide_data.get("presenter_script", "")

            out_filename = f"{title_slug}_{uid}.pptx"
            out_path = os.path.join(static_dir, out_filename)
            prs.save(out_path)
            return FileResponse(out_path, filename=out_filename, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")

        else:
            raise HTTPException(status_code=400, detail="Unsupported format. Use 'pptx' or 'txt'.")

    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))
