
from app.schemas_assessment import ExecutiveSummaryDraft
import os
from docx import Document
from docx.shared import Pt, RGBColor
from app.schemas_assessment import ToolAssessment

def render_assessment_report(subsidiary_name: str, assessment: ToolAssessment, out_path: str):
    template_path = os.path.join('static', 'reports', 'GTBank_Tool_Assessment_Template.docx')
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found at {template_path}")
        
    doc = Document(template_path)
    
    # Sort findings: CRITICAL > HIGH > MEDIUM > LOW, then by finding_number
    severity_order = {"CRITICAL": 1, "HIGH": 2, "MEDIUM": 3, "LOW": 4}
    sorted_findings = sorted(
        assessment.findings, 
        key=lambda f: (severity_order.get(f.severity, 99), f.finding_number or 999)
    )
    
    # Helper for safe text replacement
    def _replace_paragraph_text(paragraph, new_text):
        if paragraph.runs:
            for i, r in enumerate(paragraph.runs):
                if i == 0:
                    r.text = new_text
                else:
                    r.text = ""
        else:
            paragraph.text = new_text

    # Find placeholders and target insertion point
    findings_heading = None
    for p in doc.paragraphs:
        # Fill title and performed by
        if "GTBANK [SUBSIDIARY]:" in p.text:
            new_text = p.text.replace("[SUBSIDIARY]", assessment.subsidiary_name if hasattr(assessment, 'subsidiary_name') else subsidiary_name)
            new_text = new_text.replace("[TOOL NAME]", assessment.tool_display_name)
            new_text = new_text.replace("[DD/MM/YYYY]", assessment.assessment_date)
            _replace_paragraph_text(p, new_text)
        elif "Performed by:" in p.text:
            new_text = p.text.replace("[Assessor Name]", assessment.performed_by)
            _replace_paragraph_text(p, new_text)
        elif "[Tool Name]" in p.text:
            new_text = p.text.replace("[Tool Name]", assessment.tool_display_name)
            _replace_paragraph_text(p, new_text)
            
        if "Findings and Recommendations" in p.text.strip():
            findings_heading = p
            
    # Delete filler text after findings heading
    if findings_heading:
        curr = findings_heading._element.getnext()
        while curr is not None:
            nxt = curr.getnext()
            curr.getparent().remove(curr)
            curr = nxt
            
    # Insert findings
    if findings_heading:
        def insert_p_func(text="", style=None):
            p = doc.add_paragraph("")
            if style:
                p.style = style
            else:
                p.style = "Normal"
            if text:
                r = p.add_run(text)
                r.font.name = "Abadi"
                r.font.size = Pt(12)
            return p

        for i, finding in enumerate(sorted_findings):
            # Number and title
            p_title = insert_p_func("", style="Heading 2" if "Heading 2" in [s.name for s in doc.styles] else "Normal")
            run = p_title.add_run(f"{i+1}. {finding.title} - {finding.severity}")
            run.font.name = "Abadi"
            run.font.size = Pt(12)
            run.bold = True
            
            if finding.severity == "CRITICAL":
                run.font.color.rgb = RGBColor(139, 0, 0) # Dark Red
            elif finding.severity == "HIGH":
                run.font.color.rgb = RGBColor(255, 0, 0) # Red
            elif finding.severity == "MEDIUM":
                run.font.color.rgb = RGBColor(255, 165, 0) # Orange/Yellow
                
            # Impact
            p_impact = insert_p_func("", style="List Bullet" if "List Bullet" in [s.name for s in doc.styles] else "Normal")
            r1 = p_impact.add_run("Impact: ")
            r1.bold = True
            r1.font.name = "Abadi"
            r1.font.size = Pt(12)
            
            r2 = p_impact.add_run(finding.impact)
            r2.font.name = "Abadi"
            r2.font.size = Pt(12)
            
            # Recommendation
            p_rec = insert_p_func("", style="List Bullet" if "List Bullet" in [s.name for s in doc.styles] else "Normal")
            r3 = p_rec.add_run("Recommendation: ")
            r3.bold = True
            r3.font.name = "Abadi"
            r3.font.size = Pt(12)
            
            r4 = p_rec.add_run(finding.recommendation)
            r4.font.name = "Abadi"
            r4.font.size = Pt(12)
            
            insert_p_func("") # spacing
            
        # Optional Advisories
        for i, adv in enumerate(assessment.advisories):
            p_adv = insert_p_func("", style="Heading 2" if "Heading 2" in [s.name for s in doc.styles] else "Normal")
            run = p_adv.add_run(f"Advisory: {adv.topic}")
            run.font.name = "Abadi"
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(0, 102, 204) # Blue
            
            p_ctx = insert_p_func(adv.context)
            
            p_rec = insert_p_func("", style="List Bullet" if "List Bullet" in [s.name for s in doc.styles] else "Normal")
            r3 = p_rec.add_run("Recommendation: ")
            r3.bold = True
            r3.font.name = "Abadi"
            r3.font.size = Pt(12)
            
            r4 = p_rec.add_run(adv.recommendation)
            r4.font.name = "Abadi"
            r4.font.size = Pt(12)
            
            insert_p_func("")
            
    doc.save(out_path)


from app.schemas_assessment import StandaloneAssessmentDraft
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def _set_shading(p, color_hex):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    p._p.get_or_add_pPr().append(shd)

def render_standalone_report(draft: StandaloneAssessmentDraft, output_path: str):
    import os
    template_path = os.path.join(os.getcwd(), "static", "reports", "GTCO Assessment Report - Master Template.docx")
    if not os.path.exists(template_path):
        raise FileNotFoundError("Master template not found.")
        
    doc = Document(template_path)
    disable_update_fields(doc)
    
    # 1. Replace Placeholders on cover
    for p in doc.paragraphs:
        if "{ASSESSMENT NAME}" in p.text:
            p.text = p.text.replace("{ASSESSMENT NAME}", draft.assessment_name.upper())
            for r in p.runs:
                r.font.name = "Abadi"
                r.font.size = Pt(16)
                r.font.bold = True
                
        if "{SCOPE}" in p.text:
            p.text = p.text.replace("{SCOPE}", draft.scope.upper())
            for r in p.runs:
                r.font.name = "Abadi"
                r.font.size = Pt(16)
                r.font.bold = True
                
    # 2. Append Content to end of doc (assuming template is just cover + doc control, or we just append)
    doc.add_page_break()
    
    # Introduction
    p_intro_head = doc.add_paragraph("Introduction")
    p_intro_head.style = "Normal"
    for r in p_intro_head.runs:
        r.font.name = "Abadi"
        r.font.bold = True
        r.font.size = Pt(13)
        
    p_intro_body = doc.add_paragraph(draft.introduction)
    p_intro_body.style = "Normal"
    for r in p_intro_body.runs: r.font.name = "Abadi"; r.font.size = Pt(11)
    
    doc.add_paragraph("")
    
    # Scope
    p_scope_head = doc.add_paragraph("Scope of Assessment")
    p_scope_head.style = "Normal"
    for r in p_scope_head.runs:
        r.font.name = "Abadi"
        r.font.bold = True
        r.font.size = Pt(13)
        
    p_scope_intro = doc.add_paragraph(f"The scope of the {draft.assessment_name} assessment includes:")
    p_scope_intro.style = "Normal"
    for r in p_scope_intro.runs: r.font.name = "Abadi"; r.font.size = Pt(11)
    
    for item in draft.scope_items:
        p_item = doc.add_paragraph(item, style="List Bullet")
        for r in p_item.runs: r.font.name = "Abadi"; r.font.size = Pt(11)
        
    doc.add_paragraph("")
    
    # Findings
    for idx, finding in enumerate(draft.findings):
        # Finding Heading
        p_find = doc.add_paragraph()
        r_title = p_find.add_run(f"{idx+1}. {finding.title.upper()} - ")
        r_title.font.name = "Abadi"
        r_title.font.bold = True
        r_title.font.size = Pt(12)
        
        r_sev = p_find.add_run(finding.severity.upper())
        r_sev.font.name = "Abadi"
        r_sev.font.bold = True
        r_sev.font.size = Pt(12)
        # Word Native Highlight
        from docx.enum.text import WD_COLOR_INDEX
        if finding.severity.upper() == "HIGH":
            r_sev.font.highlight_color = WD_COLOR_INDEX.RED
        elif finding.severity.upper() == "MID":
            r_sev.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r_sev.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
            
        doc.add_paragraph("")
        
        def add_banner_and_list(label, items):
            p_ban = doc.add_paragraph()
            r_ban = p_ban.add_run(f" {label}")
            r_ban.font.name = "Abadi"
            r_ban.font.bold = True
            r_ban.font.size = Pt(11)
            r_ban.font.color.rgb = RGBColor(255, 255, 255)
            _set_shading(p_ban, "2E5FA2")
            
            if not items:
                p_item = doc.add_paragraph("None identified.", style="List Bullet")
                for r in p_item.runs: r.font.name = "Abadi"; r.font.size = Pt(11)
            else:
                for it in items:
                    p_item = doc.add_paragraph(it, style="List Bullet")
                    for r in p_item.runs: r.font.name = "Abadi"; r.font.size = Pt(11)
            doc.add_paragraph("")
            
        add_banner_and_list("Observations", finding.observations)
        add_banner_and_list("Impact", finding.impact)
        add_banner_and_list("Recommendations", finding.recommendations)
        
    # Conclusion
    doc.add_paragraph("")
    p_con_head = doc.add_paragraph("CONCLUSION")
    p_con_head.style = "Normal"
    for r in p_con_head.runs:
        r.font.name = "Abadi"
        r.font.bold = True
        r.font.size = Pt(13)
        
    p_con_body = doc.add_paragraph(f"This concludes the {draft.assessment_name} configuration review.")
    p_con_body.style = "Normal"
    for r in p_con_body.runs: r.font.name = "Abadi"; r.font.size = Pt(11)
    
    doc.save(output_path)


def render_executive_summary_report(draft: ExecutiveSummaryDraft, output_path: str):
    import os
    from docx import Document
    from docx.shared import Pt, RGBColor
    
    template_path = os.path.join(os.getcwd(), "static", "reports", "Executive_Summary_Template.docx")
    if not os.path.exists(template_path):
        doc = Document()
    else:
        doc = Document(template_path)
        
    # Replace standard placeholders in the template if they exist
    for p in doc.paragraphs:
        if "[BANK NAME]" in p.text:
            p.text = p.text.replace("[BANK NAME]", draft.bank_name)
        if "[COUNTRY]" in p.text:
            p.text = p.text.replace("[COUNTRY]", draft.country)
        if "[start date]" in p.text:
            p.text = p.text.replace("[start date] – [end date]", draft.assessment_period)
            
    # Append content sequentially (as a simplistic generation for now)
    doc.add_page_break()
    
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(f"Assessment Period: {draft.assessment_period}")
    doc.add_paragraph(f"Scope: {', '.join(draft.scope_domains)}")
    doc.add_paragraph(f"Overall Security Posture: {draft.overall_posture_rating} - {draft.overall_posture_summary}")
    
    doc.add_heading("Key Highlights", level=2)
    doc.add_paragraph("Highlight of Critical Observations:")
    for h in draft.critical_highlights:
        doc.add_paragraph(f"{h.domain}: {h.finding_highlight}", style='List Bullet')
        
    doc.add_paragraph("Impact:")
    for imp in draft.impact_highlights:
        doc.add_paragraph(imp, style='List Bullet')
        
    doc.add_heading("Total Observations Summary", level=2)
    table1 = doc.add_table(rows=1, cols=6)
    table1.style = 'Table Grid'
    hdr_cells = table1.rows[0].cells
    hdr_cells[0].text = 'Scope'
    hdr_cells[1].text = 'Critical'
    hdr_cells[2].text = 'High'
    hdr_cells[3].text = 'Medium'
    hdr_cells[4].text = 'Low'
    hdr_cells[5].text = 'Total'
    for row in draft.table1_rows:
        row_cells = table1.add_row().cells
        row_cells[0].text = row.domain
        row_cells[1].text = str(row.critical)
        row_cells[2].text = str(row.high)
        row_cells[3].text = str(row.medium)
        row_cells[4].text = str(row.low)
        row_cells[5].text = str(row.total)
        
    doc.add_heading("Additional Recommendations", level=2)
    for rec in draft.additional_recommendations:
        doc.add_paragraph(rec.title, style='List Bullet')
        p = doc.add_paragraph(f"Impact: {rec.impact}")
        p.paragraph_format.left_indent = Pt(36)
        p = doc.add_paragraph(f"Action: {rec.action}")
        p.paragraph_format.left_indent = Pt(36)
        
    doc.add_page_break()
    doc.add_heading("Information Security Risk Management", level=1)
    doc.add_paragraph(f"Scope: {draft.risk_scope}")
    
    table2 = doc.add_table(rows=1, cols=5)
    table2.style = 'Table Grid'
    hdr2_cells = table2.rows[0].cells
    hdr2_cells[0].text = 'Risk Rating'
    hdr2_cells[1].text = 'Critical'
    hdr2_cells[2].text = 'High'
    hdr2_cells[3].text = 'Medium'
    hdr2_cells[4].text = 'Total'
    for row in draft.table2_rows:
        row_cells = table2.add_row().cells
        row_cells[0].text = row.risk_rating
        row_cells[1].text = str(row.critical)
        row_cells[2].text = str(row.high)
        row_cells[3].text = str(row.medium)
        row_cells[4].text = str(row.total)
        
    doc.add_heading("Key Highlights (Critical & High)", level=2)
    for kf in draft.key_findings:
        doc.add_paragraph(f"{kf.finding_name}: {kf.description}", style='List Bullet')
        
    doc.add_heading("Recommendations", level=2)
    for kf in draft.key_findings:
        doc.add_paragraph(kf.recommendation_step, style='List Bullet')
        
    doc.add_page_break()
    doc.add_heading("SOC Maturity Assessment Summary", level=1)
    doc.add_paragraph(draft.soc_purpose)
    doc.add_paragraph(f"Frameworks Used: {draft.soc_frameworks}")
    doc.add_paragraph(f"Overall Maturity: {draft.soc_overall_maturity_rating} - {draft.soc_overall_maturity_rationale}")
    
    doc.add_heading("Current Scores by NIST Functions:", level=2)
    for n in draft.nist_scores:
        doc.add_paragraph(f"{n.function}: {n.score} ({n.rationale})", style='List Bullet')
        
    doc.add_paragraph(f"SOC-CMM Target: {draft.soc_cmm_target}")
    
    doc.add_heading("Key Gaps", level=2)
    for gap in draft.soc_key_gaps:
        doc.add_paragraph(gap, style='List Bullet')
        
    doc.add_heading("Recommended Roadmap", level=2)
    for yr in draft.soc_roadmap:
        doc.add_paragraph(f"{yr.year_label}: {yr.theme} - {yr.initiatives}", style='List Bullet')
        
    doc.save(output_path)
