import os
from docx import Document

def _replace_text_after_colon(paragraph, new_text):
    """Finds a colon in a paragraph and replaces everything after it."""
    if not paragraph.runs:
        return
        
    full_text = paragraph.text
    if ":" in full_text:
        label = full_text.split(":", 1)[0] + ":"
        # Clear existing runs
        for r in paragraph.runs:
            r.text = ""
        # Re-add label as bold
        label_run = paragraph.add_run(label)
        label_run.bold = True
        # Add new text
        paragraph.add_run(f" {new_text}")

def _replace_paragraph_text(paragraph, new_text):
    """Safely replaces text in a paragraph, attempting to keep first run's formatting."""
    if paragraph.runs:
        for i, r in enumerate(paragraph.runs):
            if i == 0:
                r.text = new_text
            else:
                r.text = ""
    else:
        paragraph.text = new_text

def _clear_table_data_rows(table):
    """Removes all rows from a table except the header (row 0)."""
    for i in range(len(table.rows) - 1, 0, -1):
        row = table.rows[i]
        row._element.getparent().remove(row._element)

def _delete_paragraphs_between_elements(p_start, p_end):
    """Deletes all XML elements between two paragraph elements."""
    if p_start is None or p_end is None:
        return
    curr = p_start._element.getnext()
    while curr is not None and curr is not p_end._element:
        nxt = curr.getnext()
        curr.getparent().remove(curr)
        curr = nxt

def _delete_element_and_following(p_start, p_end):
    """Deletes p_start and all elements up to (but not including) p_end."""
    if p_start is None or p_end is None:
        return
    curr = p_start._element
    while curr is not None and curr is not p_end._element:
        nxt = curr.getnext()
        curr.getparent().remove(curr)
        curr = nxt

def render_health_check_report(document_model: dict, output_path: str):
    template_path = os.path.abspath(os.path.join(os.path.dirname(__file__), 'health_check_template.docx'))
    doc = Document(template_path)
    
    # 1. Cover Page Title
    for p in doc.paragraphs[:20]:
        if "HEALTH CHECK REMEDIATION REPORT" in p.text.upper() or "Q1 WAF" in p.text.upper():
            _replace_paragraph_text(p, document_model.get('report_title', 'Health Check Remediation Report').upper())
            p.runs[0].bold = True
            break
            
    # 2. Map all headings to their paragraph objects
    h_elements = {}
    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith('Report Summary'): h_elements['summary'] = p
        elif text.startswith('Introduction'): h_elements['intro'] = p
        elif text.startswith('Observations and Remediations'): h_elements['obs'] = p
        elif text.startswith('Summary of Remediation Status'): h_elements['status'] = p
        elif text.startswith('Limitation'): h_elements['limitation'] = p
        elif text in ['Conclusion', 'Conclusions']: h_elements['conclusion'] = p
        elif text == 'Appendices': h_elements['appendices'] = p
        elif text == 'References': h_elements['references'] = p
        elif text in ['Sign off', 'Sign Off']: h_elements['signoff'] = p

    # --- DELETE LEGACY CONTENT & INSERT NEW CONTENT ---
    
    # References to Signoff
    if 'references' in h_elements and 'signoff' in h_elements:
        _delete_paragraphs_between_elements(h_elements['references'], h_elements['signoff'])
        h_elements['signoff'].insert_paragraph_before("\n".join(document_model.get('references', ['None'])))
    
    # Appendices to References
    if 'appendices' in h_elements and 'references' in h_elements:
        _delete_paragraphs_between_elements(h_elements['appendices'], h_elements['references'])
        h_elements['references'].insert_paragraph_before("\n".join(document_model.get('appendices', ['None'])))
        
    # Conclusion to Appendices
    if 'conclusion' in h_elements and 'appendices' in h_elements:
        _delete_paragraphs_between_elements(h_elements['conclusion'], h_elements['appendices'])
        h_elements['appendices'].insert_paragraph_before(document_model.get('conclusion', ''))
        
    # Limitations to Conclusion (We completely nuke Limitations if it exists)
    if 'limitation' in h_elements and 'conclusion' in h_elements:
        _delete_element_and_following(h_elements['limitation'], h_elements['conclusion'])

    # Observations to Status
    if 'obs' in h_elements and 'status' in h_elements:
        _delete_paragraphs_between_elements(h_elements['obs'], h_elements['status'])
        
        insert_target = h_elements['status']
        observations = document_model.get('observations', [])
        for i, obs in enumerate(observations):
            insert_target.insert_paragraph_before('')
            
            p_title = insert_target.insert_paragraph_before('')
            t_label = p_title.add_run(f"{i + 1}.     {obs.get('title', '')}")
            t_label.bold = True
            
            p_obs = insert_target.insert_paragraph_before('')
            o_label = p_obs.add_run("Observation: ")
            o_label.bold = True
            p_obs.add_run(obs.get('observation', ''))
            
            p_rem = insert_target.insert_paragraph_before('')
            r_label = p_rem.add_run(f"{obs.get('remediation_label', 'Remediation Action')}: ")
            r_label.bold = True
            p_rem.add_run(obs.get('remediation_text', ''))
            
            insert_target.insert_paragraph_before('')
            
    # Intro to Observations
    if 'intro' in h_elements and 'obs' in h_elements:
        _delete_paragraphs_between_elements(h_elements['intro'], h_elements['obs'])
        h_elements['obs'].insert_paragraph_before(document_model.get('introduction', ''))
        
    # --- UPDATE SUMMARY METADATA ---
    if 'summary' in h_elements and 'intro' in h_elements:
        # Traverse paragraphs between summary and intro
        curr = h_elements['summary']._element.getnext()
        while curr is not None and curr is not h_elements['intro']._element:
            # Reconstruct paragraph wrapper to use text methods
            from docx.text.paragraph import Paragraph
            if curr.tag.endswith('p'):
                p = Paragraph(curr, doc._body)
                text = p.text.strip()
                if text.startswith('Report Title:'):
                    _replace_text_after_colon(p, document_model.get('report_title', ''))
                elif text.startswith('Date:'):
                    _replace_text_after_colon(p, document_model.get('report_date', ''))
                elif text.startswith('Prepared By:'):
                    _replace_text_after_colon(p, document_model.get('prepared_by', ''))
                elif text.startswith('Reviewed By:'):
                    _replace_text_after_colon(p, document_model.get('reviewed_by', ''))
                elif text.startswith('Organization/Unit:'):
                    _replace_text_after_colon(p, document_model.get('organization_unit', 'Security Monitoring and Threat Intelligence'))
            curr = curr.getnext()

    # --- UPDATE TABLES ---
    if doc.tables:
        status_table = doc.tables[0]
        _clear_table_data_rows(status_table)
        
        observations = document_model.get('observations', [])
        for obs in observations:
            row = status_table.add_row()
            row.cells[0].text = obs.get('title', '')
            row.cells[1].text = obs.get('status', 'Pending')
            
    if len(doc.tables) > 1:
        sign_table = doc.tables[1]
        if len(sign_table.rows) > 0 and len(sign_table.columns) > 1:
            rev_cell = sign_table.cell(0, 0)
            rev_cell.text = f"Reviewed by: {document_model.get('reviewed_by', '')}\n____________________________\nPosition: "
            
            app_cell = sign_table.cell(0, 1)
            app_cell.text = f"Approved by: \n____________________________\nPosition: "

    doc.save(output_path)
    return output_path
